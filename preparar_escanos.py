#!/usr/bin/env python3
"""
Cuántos concejales urbanos y rurales elige cada cantón.

El tablero necesita saberlo para no preguntárselo al usuario y para dejar el
número FIJO en la pantalla de asignación.

Fuente
  · Escanos_2023-consolidado.xlsx, entregado por el usuario, con cuatro hojas:
      CONCEJALES_URBANOS          un cantón por fila
      CONCEJALES_URBANOS_CIRCUNS  los cantones grandes, partidos en
                                  circunscripciones urbanas (una fila cada una)
      CONCEJALES_RURALES          un cantón por fila
      JUNTAS_PARROQUIALES         vocales por parroquia

  Los cantones con circunscripciones NO están en la primera hoja: las dos se
  suman. Para esos cantones se guarda además el detalle, porque la elección
  real se reparte dentro de cada circunscripción y no en el cantón entero.

  La hoja de juntas parroquiales no se usa: en este tablero los escaños solo
  se aplican a concejales.

Nombres
  Los de la hoja vienen con tildes y con abreviaturas del CNE. Se normalizan
  igual que en js/data.js y se traducen con las mismas tablas de alias que ya
  usaba el tablero; al final se comprueba que TODOS los cantones existan en
  data/electores.json, que es la lista que maneja el tablero. Si alguno no
  cuadra, el script se detiene y lo dice: más vale no publicar un escaño
  colgado de una clave que nadie va a encontrar.

Salida: data/escanos.json
    {"anio": 2023, "fuente": "...xlsx",
     "escanos": {"CONCEJALES_URBANOS": {"PICHINCHA|QUITO": 15, ...},
                 "CONCEJALES_RURALES": {"PICHINCHA|QUITO": 6, ...}},
     "circunscripciones": {"CONCEJALES_URBANOS": {"PICHINCHA|QUITO": [5,5,5], ...}}}

Uso
    python preparar_escanos.py RUTA/Escanos_2023-consolidado.xlsx

Requiere openpyxl (pip install openpyxl).
"""
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

BASE = Path(__file__).parent
DESTINO = BASE / 'data' / 'escanos.json'
CANTONES = BASE / 'data' / 'electores.json'

ANIO = 2023
HOJAS = {
    'CONCEJALES_URBANOS': ('CONCEJALES_URBANOS', 'CONCEJALES_URBANOS_CIRCUNS'),
    'CONCEJALES_RURALES': ('CONCEJALES_RURALES',),
}

ALIAS_PROV = {
    'STO DGO TSACHILAS': 'SANTO DOMINGO DE LOS TSACHILAS',
    'SANTO DOMINGO': 'SANTO DOMINGO DE LOS TSACHILAS',
    'STO DOMINGO': 'SANTO DOMINGO DE LOS TSACHILAS',
    'SANTO DOMINGO TSACHILAS': 'SANTO DOMINGO DE LOS TSACHILAS',
}
ALIAS_CANT = {
    'A BAQUERIZO MORENO': 'ALFREDO BAQUERIZO MORENO',
    'CRNL MARCELINO MARIDUENAS': 'CRNEL MARCELINO MARIDUENA',
    'CORONEL MARCELINO MARIDUENA': 'CRNEL MARCELINO MARIDUENA',
    'EL EMPALME': 'EMPALME', 'GRAL A ELIZALDE': 'GNRAL ANTONIO ELIZALDE',
    'GENERAL ANTONIO ELIZALDE': 'GNRAL ANTONIO ELIZALDE',
    'NOBOL PIEDRAHITA': 'NOBOL', 'YAGUACHI': 'SAN JACINTO DE YAGUACHI',
    'C J AROSEMENA TOLA': 'CARLOS JULIO AROSEMENA TOLA',
    'FCO DE ORELLANA': 'ORELLANA', 'FRANCISCO DE ORELLANA': 'ORELLANA',
    'JOYA DE LOS SACHAS': 'LA JOYA DE LOS SACHAS',
    'PUEBLO VIEJO': 'PUEBLOVIEJO', 'RIO VERDE': 'RIOVERDE',
    'URCUQUI': 'SAN MIGUEL DE URCUQUI', 'BANOS': 'BANOS DE AGUA SANTA',
    'PELILEO': 'SAN PEDRO DE PELILEO', 'PILLARO': 'SANTIAGO DE PILLARO',
    'YANZATZA': 'YANTZAZA',
}


def norm(s):
    """Réplica de normalizarTexto() de js/data.js."""
    s = str(s if s is not None else '').upper().replace('Ñ', '\x01')
    s = ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
    s = s.replace('\x01', 'N')
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()


def clave(prov, cant):
    p = ALIAS_PROV.get(norm(prov), norm(prov))
    c = ALIAS_CANT.get(norm(cant), norm(cant))
    return f'{p}|{c}'


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    ruta = Path(sys.argv[1])
    print(f'Leyendo {ruta.name}…')
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    conocidos = set(json.loads(CANTONES.read_text(encoding='utf-8'))['cantones'])
    print(f'  el tablero conoce {len(conocidos)} cantones')

    escanos, circuns, sueltos = {}, {}, []
    for ident, hojas in HOJAS.items():
        cuenta = defaultdict(int)
        detalle = defaultdict(list)
        for hoja in hojas:
            if hoja not in wb.sheetnames:
                sys.exit(f'Falta la hoja «{hoja}» en el Excel.')
            ws = wb[hoja]
            filas = list(ws.iter_rows(values_only=True))
            cab = [norm(x) for x in filas[0]]
            iProv, iCant = cab.index('PROVINCIA'), cab.index('CANTON')
            iEsc = cab.index('ESCANOS')
            iCirc = cab.index('CIRCUNSCRIPCION') if 'CIRCUNSCRIPCION' in cab else None
            n = 0
            for f in filas[1:]:
                if f[iProv] is None or f[iEsc] is None:
                    continue
                k = clave(f[iProv], f[iCant])
                cuenta[k] += int(f[iEsc])
                if iCirc is not None:
                    detalle[k].append(int(f[iEsc]))
                if k not in conocidos:
                    sueltos.append((hoja, f[iProv], f[iCant], k))
                n += 1
            print(f'  {hoja}: {n} filas')
        escanos[ident] = dict(sorted(cuenta.items()))
        if detalle:
            circuns[ident] = dict(sorted(detalle.items()))

    if sueltos:
        print('\n✗ Cantones del Excel que el tablero no reconoce:')
        for hoja, p, c, k in sueltos:
            print(f'    {hoja}: «{p} / {c}» → {k}')
        sys.exit('Añade el alias que falte en ALIAS_PROV/ALIAS_CANT y vuelve a correrlo.')

    # ── Comparación con lo que había, para no cambiar a ciegas ──────────────
    if DESTINO.exists():
        viejo = json.loads(DESTINO.read_text(encoding='utf-8'))
        print(f'\n  Frente al archivo anterior (de {viejo.get("anio")}):')
        for ident in HOJAS:
            a, b = viejo['escanos'].get(ident, {}), escanos[ident]
            iguales = sum(1 for k in b if a.get(k) == b[k])
            nuevos = [k for k in b if k not in a]
            distintos = {k: (a[k], b[k]) for k in b if k in a and a[k] != b[k]}
            print(f'    {ident}: {len(b)} cantones · {iguales} iguales · '
                  f'{len(distintos)} cambian · {len(nuevos)} nuevos')
            for k, (x, y) in list(sorted(distintos.items()))[:6]:
                print(f'       {k:<46} {x} → {y}')
            if len(distintos) > 6:
                print(f'       … y {len(distintos) - 6} más')

    salida = {'anio': ANIO, 'fuente': ruta.name, 'escanos': escanos,
              'circunscripciones': circuns}
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    DESTINO.write_text(json.dumps(salida, ensure_ascii=False, separators=(',', ':')),
                       encoding='utf-8')
    print(f'\n✓ {DESTINO}  ({DESTINO.stat().st_size / 1024:.0f} KB)')

    for ident in HOJAS:
        v = escanos[ident]
        print(f'  {ident}: {len(v)} cantones, {sum(v.values())} escaños '
              f'(de {min(v.values())} a {max(v.values())} por cantón)')
    if circuns.get('CONCEJALES_URBANOS'):
        c = circuns['CONCEJALES_URBANOS']
        print(f'  {len(c)} cantones eligen sus concejales urbanos por circunscripción:')
        for k, v in sorted(c.items()):
            print(f'     {k:<46} {" + ".join(map(str, v))} = {sum(v)}')

    print('\n  Para comprobar a ojo (urbanos + rurales = total del concejo):')
    for p, c in [('PICHINCHA', 'QUITO'), ('GUAYAS', 'GUAYAQUIL'), ('AZUAY', 'CUENCA'),
                 ('MANABI', 'PORTOVIEJO'), ('COTOPAXI', 'LA MANA'), ('MANABI', 'PEDERNALES')]:
        k = f'{p}|{c}'
        u = escanos['CONCEJALES_URBANOS'].get(k, 0)
        r = escanos['CONCEJALES_RURALES'].get(k, 0)
        print(f'     {c:<12} {u:>3} urbanos + {r:>2} rurales = {u + r:>3}')


if __name__ == '__main__':
    main()
