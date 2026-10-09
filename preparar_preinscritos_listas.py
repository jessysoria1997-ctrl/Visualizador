#!/usr/bin/env python3
"""
Añade a data/preinscritos.xlsx las dignidades que se eligen por listas:
concejales (urbanos, rurales y urbanos por circunscripción) y vocales de
juntas parroquiales rurales.

Fuente
  · Reporte_compacto_Pre_inscritos_Estadistica.xls   una candidatura por fila,
    con su provincia, cantón y —en los vocales— su parroquia

Qué hace
  · Deja intactas las filas de prefectos y alcaldes: las de concejales y
    vocales se reemplazan por completo (se puede volver a ejecutar sin que
    queden duplicados).
  · Cada columna del reporte va a su columna de la base:
        DIGNIDAD                 → DIGNIDAD   (tal cual: el tablero reconoce
                                               «CONCEJAL» y «VOCAL» en el texto)
        PROVINCIA, CANTÓN        → PROVINCIA, CANTON
        PARROQUIA                → PARROQUIA  (solo los vocales la traen)
        LISTA                    → LISTA
        PARTIDO                  → PARTIDO
        CÉDULA, NOMBRE           → CEDULA, NOMBRE
        ORDEN_CANDIDATO          → ORDEN
        CEDULA_SUPLENTE          → CEDULA_SUPLENTE
        NOM_CANDIDATO_SUPLENTE   → NOMBRE_SUPLENTE
        CIRCUNS. CANTONAL        → CIRCUNSCRIPCION
    Solo se quitan espacios sobrantes; ningún valor se cambia ni se inventa.
  · Informa de las parroquias que no cruzan con la capa parroquial: sus
    candidaturas cuentan en el cantón y la provincia, pero no se pintan en el
    mapa parroquial.

Uso
    python preparar_preinscritos_listas.py RUTA/Reporte_compacto_Pre_inscritos_Estadistica.xls

Requiere pandas, openpyxl y xlrd (pip install xlrd) para leer el .xls.
"""
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd

BASE = Path(__file__).parent
BASE_XLSX = BASE / 'data' / 'preinscritos.xlsx'

# Dignidades que trae este reporte (texto contenido en DIGNIDAD).
CLAVES = ['CONCEJAL', 'VOCAL']

# Columna de la base ← primera columna del reporte que exista.
COLUMNAS = {
    'DIGNIDAD': ['DIGNIDAD'],
    'PROVINCIA': ['PROVINCIA'],
    'CANTON': ['CANTÓN', 'CANTON'],
    'PARROQUIA': ['PARROQUIA'],
    'LISTA': ['LISTA'],
    'PARTIDO': ['PARTIDO', 'ORGANIZACIÓN POLÍTICA'],
    'CEDULA': ['CÉDULA', 'CEDULA'],
    'NOMBRE': ['NOMBRE'],
    'ORDEN': ['ORDEN_CANDIDATO'],
    'CEDULA_SUPLENTE': ['CEDULA_SUPLENTE'],
    'NOMBRE_SUPLENTE': ['NOM_CANDIDATO_SUPLENTE'],
    'CIRCUNSCRIPCION': ['CIRCUNS. CANTONAL'],
}

ALIAS_PROV = {
    'STO DGO TSACHILAS': 'SANTO DOMINGO DE LOS TSACHILAS',
    'SANTO DOMINGO': 'SANTO DOMINGO DE LOS TSACHILAS',
    'STO DOMINGO': 'SANTO DOMINGO DE LOS TSACHILAS',
    'SANTO DOMINGO TSACHILAS': 'SANTO DOMINGO DE LOS TSACHILAS',
}
ALIAS_CANT = {
    'A BAQUERIZO MORENO': 'ALFREDO BAQUERIZO MORENO',
    'CRNL MARCELINO MARIDUENAS': 'CRNEL MARCELINO MARIDUENA',
    'EL EMPALME': 'EMPALME', 'GRAL A ELIZALDE': 'GNRAL ANTONIO ELIZALDE',
    'NOBOL PIEDRAHITA': 'NOBOL', 'YAGUACHI': 'SAN JACINTO DE YAGUACHI',
    'C J AROSEMENA TOLA': 'CARLOS JULIO AROSEMENA TOLA',
    'FCO DE ORELLANA': 'ORELLANA', 'JOYA DE LOS SACHAS': 'LA JOYA DE LOS SACHAS',
    'PUEBLO VIEJO': 'PUEBLOVIEJO', 'RIO VERDE': 'RIOVERDE',
    'URCUQUI': 'SAN MIGUEL DE URCUQUI', 'BANOS': 'BANOS DE AGUA SANTA',
    'PELILEO': 'SAN PEDRO DE PELILEO', 'PILLARO': 'SANTIAGO DE PILLARO',
    'YANZATZA': 'YANTZAZA',
}


def norm(s):
    """Réplica de normalizarTexto() de js/data.js."""
    s = str(s if s is not None else '').upper().replace('Ñ', '\x01')
    s = ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
    s = s.replace('\x01', 'N')
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()


def texto(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    t = re.sub(r'\s+', ' ', str(x)).strip()
    return t or None


def revisar_parroquias(filas):
    """Informe: parroquias del reporte que no tienen polígono en la capa."""
    capa = json.loads((BASE / 'data' / 'parroquias.geojson').read_text(encoding='utf-8'))
    por_cant = defaultdict(list)
    for f in capa['features']:
        p = f['properties']
        nombres = {norm(p['dpa_despar'])} | {norm(x) for x in str(p.get('nombres_cne') or '').split('|') if norm(x)}
        por_cant[(norm(p['dpa_despro']), norm(p['dpa_descan']))].append(nombres)

    def cruza(prov, cant, parr):
        ps = por_cant.get((prov, cant), [])
        return any(parr in n or parr.replace(' ', '') in {x.replace(' ', '') for x in n} for n in ps)

    sin, con = defaultdict(int), set()
    for f in filas:
        if not f['PARROQUIA']:
            continue
        prov = ALIAS_PROV.get(norm(f['PROVINCIA']), norm(f['PROVINCIA']))
        cant = ALIAS_CANT.get(norm(f['CANTON']), norm(f['CANTON']))
        parr = norm(f['PARROQUIA'])
        if cruza(prov, cant, parr):
            con.add((prov, cant, parr))
        else:
            sin[(prov, cant, parr)] += 1
    print(f'  parroquias con polígono en el mapa: {len(con)}')
    if sin:
        print(f'  sin polígono ({len(sin)}; cuentan en su cantón y provincia):')
        for (p, c, r), n in sorted(sin.items()):
            print(f'     {p} · {c} · {r}: {n} candidaturas')


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    rep = pd.read_excel(sys.argv[1], dtype=str)
    rep.columns = [str(c).strip() for c in rep.columns]
    cols = {}
    for destino, posibles in COLUMNAS.items():
        hallada = next((c for c in posibles if c in rep.columns), None)
        if not hallada:
            sys.exit(f'Al reporte le falta la columna {posibles[0]}')
        cols[destino] = hallada

    rep = rep[rep[cols['DIGNIDAD']].notna()]          # la última fila es el pie
    otras = rep[~rep[cols['DIGNIDAD']].map(norm).str.contains('|'.join(CLAVES))]
    if len(otras):
        sys.exit(f'{len(otras)} filas no son de concejales ni de vocales: '
                 f'{sorted(otras[cols["DIGNIDAD"]].unique())}')

    filas = [{d: texto(r[cols[d]]) for d in COLUMNAS} for _, r in rep.iterrows()]

    wb = openpyxl.load_workbook(BASE_XLSX)
    ws = wb.active
    cab = [c.value for c in ws[1]]
    for d in COLUMNAS:
        if d not in cab:
            ws.cell(row=1, column=len(cab) + 1, value=d)
            cab.append(d)

    # Concejales y vocales de una ejecución anterior: siempre al final.
    idig = cab.index('DIGNIDAD')
    previas = [i for i, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
               if any(c in norm(fila[idig]) for c in CLAVES)]
    if previas:
        if previas != list(range(previas[0], ws.max_row + 1)):
            sys.exit('La base tiene concejales o vocales intercalados con otras filas: revísala a mano.')
        ws.delete_rows(previas[0], len(previas))
    originales = ws.max_row - 1

    for f in filas:
        ws.append([f.get(c) for c in cab])
    wb.save(BASE_XLSX)

    # ── Comprobación ────────────────────────────────────────────────────────
    nueva = pd.read_excel(BASE_XLSX, dtype=str)
    print(f'✓ {BASE_XLSX}')
    print(f'  {originales} filas de otras dignidades (sin cambios) + {len(filas)} de este reporte '
          f'= {len(nueva)}')
    print('  ' + ', '.join(f'{k}: {v}' for k, v in nueva['DIGNIDAD'].value_counts().items()))
    if previas:
        print(f'  (se reemplazaron {len(previas)} filas de una ejecución anterior)')
    revisar_parroquias(filas)


if __name__ == '__main__':
    main()
