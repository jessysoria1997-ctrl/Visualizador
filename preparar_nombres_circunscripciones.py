#!/usr/bin/env python3
"""
Nombres oficiales de las circunscripciones urbanas, tomados del Excel de
escaños que entregó el cliente, y cruce con los datos del tablero.

Por qué hace falta
  data/circunscripciones.json sale del distributivo del CNE de 2026 y numera
  las circunscripciones («Circunscripción Urbana 1, 2, 3»). El cliente pide
  que la sección Asignación use los nombres EXACTOS de su Excel
  (Escanos_2023-consolidado.xlsx, hoja CONCEJALES_URBANOS_CIRCUNS), que en
  Quito no son números sino «CIRCUNSCRIPCION URBANA NORTE / CENTRO / SUR».

Cómo se empareja cada circunscripción con su nombre
  1. Si el nombre del Excel termina en un número («CIRCUNSCRIPCION URBANA 2»),
     se empareja con la circunscripción de ese mismo número. Es literal.
  2. Si el nombre del Excel es de orientación (NORTE, CENTRO, SUR…) y el
     cantón tiene el mismo número de circunscripciones en los dos archivos,
     se empareja por UBICACIÓN: el centro geográfico de las parroquias de cada
     circunscripción (capa data/parroquias.geojson). La más al norte es la
     NORTE, la más al sur la SUR y la que queda, la CENTRO. El script imprime
     las latitudes para que se pueda comprobar.
  3. Si el Excel no trae la circunscripción —es nueva para 2027—, se usa el
     nombre literal del registro de inscripción (data/preinscritos.xlsx,
     columna CIRCUNSCRIPCION). No se inventa ninguno.

  El nombre se guarda tal cual: mayúsculas, sin añadir ni quitar tildes, en el
  mismo orden de palabras. Los campos que ya tenía el archivo no se tocan.

Qué se añade a cada circunscripción de data/circunscripciones.json
  nombreOficial   el nombre que muestra el tablero (literal)
  fuenteNombre    'excel' o 'registro'
  emparejado      'numero', 'ubicacion' o 'registro'
  nombreRegistro  el nombre literal en el registro de inscripción 2027
  escanosExcel    los escaños de esa circunscripción en el Excel (2023), o null
y a cada cantón un bloque 'excel' con lo que no cuadra entre los archivos.

Uso
    python preparar_nombres_circunscripciones.py RUTA/Escanos_2023-consolidado.xlsx

Correrlo DESPUÉS de preparar_circunscripciones.py (que reescribe el archivo).
Requiere openpyxl (pip install openpyxl).
"""
import json
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

BASE = Path(__file__).parent
CIRCUNS = BASE / 'data' / 'circunscripciones.json'
PARROQUIAS = BASE / 'data' / 'parroquias.geojson'
REGISTRO = BASE / 'data' / 'preinscritos.xlsx'
HOJA = 'CONCEJALES_URBANOS_CIRCUNS'

ALIAS_PROV = {
    'STO DGO TSACHILAS': 'SANTO DOMINGO DE LOS TSACHILAS',
    'SANTO DOMINGO': 'SANTO DOMINGO DE LOS TSACHILAS',
    'STO DOMINGO': 'SANTO DOMINGO DE LOS TSACHILAS',
    'SANTO DOMINGO TSACHILAS': 'SANTO DOMINGO DE LOS TSACHILAS',
}
ALIAS_CANT = {
    'YAGUACHI': 'SAN JACINTO DE YAGUACHI', 'FCO DE ORELLANA': 'ORELLANA',
    'JOYA DE LOS SACHAS': 'LA JOYA DE LOS SACHAS', 'RIO VERDE': 'RIOVERDE',
}
# Orientaciones que puede llevar un nombre, con el eje y el sentido en que se
# ordenan: la NORTE es la de mayor latitud, la ESTE la de mayor longitud.
ORIENTACION = {'NORTE': ('lat', 1), 'SUR': ('lat', -1), 'ESTE': ('lon', 1),
               'ORIENTE': ('lon', 1), 'OESTE': ('lon', -1), 'OCCIDENTE': ('lon', -1)}


def norm(s):
    """Réplica de normalizarTexto() de js/data.js."""
    s = str(s if s is not None else '').upper().replace('Ñ', '\x01')
    s = ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
    s = s.replace('\x01', 'N')
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()


def clave(prov, cant):
    p = ALIAS_PROV.get(norm(prov), norm(prov))
    c = ALIAS_CANT.get(norm(cant), norm(cant))
    return f'{p}|{c}'


def numero(nombre):
    m = re.search(r'(\d+)\s*$', str(nombre or ''))
    return int(m.group(1)) if m else None


def leer_excel(ruta):
    """{clave: [(nombre literal, escaños)]} de la hoja de circunscripciones."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    if HOJA not in wb.sheetnames:
        sys.exit(f'Falta la hoja «{HOJA}» en el Excel.')
    filas = list(wb[HOJA].iter_rows(values_only=True))
    cab = [norm(x) for x in filas[0]]
    iP, iC = cab.index('PROVINCIA'), cab.index('CANTON')
    iN, iE = cab.index('CIRCUNSCRIPCION'), cab.index('ESCANOS')
    out = defaultdict(list)
    for f in filas[1:]:
        if f[iP] is None or f[iN] is None:
            continue                       # la fila de total, sin cantón
        out[clave(f[iP], f[iC])].append((str(f[iN]).strip(), int(f[iE])))
    return out


def leer_registro():
    """{clave: {n: (nombre literal, escaños según los candidatos inscritos)}}."""
    if not REGISTRO.exists():
        return {}
    wb = openpyxl.load_workbook(REGISTRO, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = ws.iter_rows(values_only=True)
    cab = [norm(x) for x in next(filas)]
    iD, iP, iC = cab.index('DIGNIDAD'), cab.index('PROVINCIA'), cab.index('CANTON')
    iN, iL, iPa = cab.index('CIRCUNSCRIPCION'), cab.index('LISTA'), cab.index('PARTIDO')
    iNo = cab.index('NOMBRE')
    nombres, cuenta = {}, defaultdict(int)
    for f in filas:
        if not f or 'CIRCUNSCRIPCION' not in norm(f[iD]) or not f[iN]:
            continue
        k, n = clave(f[iP], f[iC]), numero(f[iN])
        if n is None:
            continue
        nombres.setdefault(k, {})[n] = str(f[iN]).strip()
        if f[iNo]:
            cuenta[(k, n, f[iL], norm(f[iPa]))] += 1
    escanos = defaultdict(int)
    for (k, n, _, _), c in cuenta.items():
        escanos[(k, n)] = max(escanos[(k, n)], c)
    return {k: {n: (v, escanos.get((k, n))) for n, v in d.items()} for k, d in nombres.items()}


def centro(geom):
    """Centro de masa (área) de un Polygon o MultiPolygon, en lon/lat."""
    polys = [geom['coordinates']] if geom['type'] == 'Polygon' else geom['coordinates']
    A = X = Y = 0.0
    for poly in polys:
        anillo = poly[0]
        for (x1, y1), (x2, y2) in zip(anillo, anillo[1:]):
            c = x1 * y2 - x2 * y1
            A += c
            X += (x1 + x2) * c
            Y += (y1 + y2) * c
    if not A:
        return None, 0.0
    return (X / (3 * A), Y / (3 * A)), abs(A) / 2


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    ruta = Path(sys.argv[1])
    print(f'Leyendo {ruta.name}…')
    excel = leer_excel(ruta)
    registro = leer_registro()
    ci = json.loads(CIRCUNS.read_text(encoding='utf-8'))
    capa = json.loads(PARROQUIAS.read_text(encoding='utf-8'))['features']

    # Parroquias de la capa por cantón, con su centro y su área.
    porCanton = defaultdict(dict)
    for f in capa:
        p = f['properties']
        k = clave(p['dpa_despro'].replace('Ð', 'Ñ'), p['dpa_descan'].replace('Ð', 'Ñ'))
        c, a = centro(f['geometry'])
        if c is None:
            continue
        nombres = {norm(p['dpa_despar'])} | {norm(x) for x in re.split(r'[|;,]', p.get('nombres_cne') or '') if x.strip()}
        for n in nombres:
            porCanton[k][n] = (c, a)
            porCanton[k][n.replace(' ', '')] = (c, a)

    def centro_circ(k, parroquias):
        sx = sy = sa = 0.0
        for nombre in parroquias:
            n = norm(nombre)
            g = porCanton[k].get(n) or porCanton[k].get(n.replace(' ', ''))
            if not g:
                continue
            (x, y), a = g
            sx += x * a; sy += y * a; sa += a
        return (sx / sa, sy / sa) if sa else None

    informe = []
    for k, cant in sorted(ci['cantones'].items()):
        urbanas = cant['urbanas']
        del_excel = excel.get(k, [])
        del_registro = registro.get(k, {})
        usados = set()
        # 1 · Por número
        for u in urbanas:
            u.pop('nombreOficial', None)
            reg = del_registro.get(u['n'])
            u['nombreRegistro'] = reg[0] if reg else None
            for i, (nombre, esc) in enumerate(del_excel):
                if i not in usados and numero(nombre) == u['n']:
                    u.update(nombreOficial=nombre, fuenteNombre='excel', emparejado='numero', escanosExcel=esc)
                    usados.add(i)
                    break
        # 2 · Por ubicación (nombres de orientación, mismo número de circunscripciones)
        orient = [(i, nombre, esc) for i, (nombre, esc) in enumerate(del_excel)
                  if i not in usados and numero(nombre) is None]
        pendientes = [u for u in urbanas if 'nombreOficial' not in u]
        if orient and len(orient) == len(pendientes) == len(del_excel):
            centros = {u['n']: centro_circ(k, u['parroquias']) for u in pendientes}
            if all(centros.values()):
                libres = list(pendientes)
                asignados = {}
                # Primero los extremos (NORTE/SUR/ESTE/OESTE), luego el CENTRO.
                for i, nombre, esc in sorted(orient, key=lambda t: norm(t[1]).split()[-1] not in ORIENTACION):
                    palabra = norm(nombre).split()[-1]
                    if palabra in ORIENTACION:
                        eje, signo = ORIENTACION[palabra]
                        j = 1 if eje == 'lat' else 0
                        u = max(libres, key=lambda u: signo * centros[u['n']][j])
                    elif len(libres) == 1:
                        u = libres[0]
                    else:
                        continue
                    libres.remove(u)
                    asignados[u['n']] = (nombre, esc)
                if not libres:
                    for u in pendientes:
                        nombre, esc = asignados[u['n']]
                        u.update(nombreOficial=nombre, fuenteNombre='excel', emparejado='ubicacion',
                                 escanosExcel=esc,
                                 centro=[round(centros[u['n']][1], 5), round(centros[u['n']][0], 5)])
                    informe.append(f'{k}: emparejadas por ubicación → ' + ', '.join(
                        f"{u['n']} = {u['nombreOficial']} (lat {centros[u['n']][1]:.4f})" for u in pendientes))
        # 3 · Del registro de inscripción
        for u in urbanas:
            if 'nombreOficial' not in u:
                if not u['nombreRegistro']:
                    sys.exit(f'{k}: la circunscripción {u["n"]} no está en el Excel ni en el registro.')
                u.update(nombreOficial=u['nombreRegistro'], fuenteNombre='registro',
                         emparejado='registro', escanosExcel=None)

        # Lo que no cuadra entre los archivos, guardado en el propio cantón
        avisos = []
        if not del_excel:
            avisos.append(f'El Excel no trae circunscripciones en este cantón; el registro de '
                          f'inscripción 2027 trae {len(urbanas)}. Sus nombres son los del registro.')
        elif len(del_excel) != len(urbanas):
            avisos.append(f'El Excel trae {len(del_excel)} circunscripciones y el registro de '
                          f'inscripción 2027, {len(urbanas)}. Las que no están en el Excel llevan '
                          f'el nombre del registro.')
        for u in urbanas:
            reg = del_registro.get(u['n'])
            if u['escanosExcel'] is not None and reg and reg[1] and reg[1] != u['escanosExcel']:
                avisos.append(f"{u['nombreOficial']}: el Excel de escaños trae {u['escanosExcel']} y el "
                              f"registro de inscripción 2027, {reg[1]} candidatos por lista; se reparten "
                              f"los del registro.")
        compartidas = sorted({p for u in urbanas for p in u['parroquias']
                              if sum(p in v['parroquias'] for v in urbanas) > 1})
        cant['excel'] = {'circunscripciones': len(del_excel), 'avisos': avisos,
                         'parroquiasCompartidas': compartidas}
        for a in avisos:
            informe.append(f'{k}: {a}')

    sobran = sorted(set(excel) - set(ci['cantones']))
    for k in sobran:
        informe.append(f'{k}: el Excel trae circunscripciones pero el distributivo 2026 no.')

    ci['nombres'] = {'fuente': ruta.name, 'hoja': HOJA,
                     'regla': 'Nombre literal del Excel; si no está, el del registro de inscripción 2027.'}
    CIRCUNS.write_text(json.dumps(ci, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    print(f'\n✓ {CIRCUNS}  ({CIRCUNS.stat().st_size / 1024:.0f} KB)\n')
    for k, cant in sorted(ci['cantones'].items()):
        print(f'  {k}')
        for u in cant['urbanas']:
            print(f"     {u['n']}  {u['nombreOficial']:<34} {u['fuenteNombre']:<9} ({u['emparejado']})")
    print('\n  Lo que no cuadra entre los archivos:')
    for linea in informe:
        print('   · ' + linea)


if __name__ == '__main__':
    main()
